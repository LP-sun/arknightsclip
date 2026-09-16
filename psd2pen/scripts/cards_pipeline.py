"""JSON-driven compiler for the existing Pencil template; never opens .pen files."""
import argparse
import hashlib
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
YS = [187,401,615,829,101,315,529,743]
LEFT = ['DLOWp','B8gA3U','MxQrR','bFlti','QEG16','Q2ZirJ','coU5m','RB169','Q3ckUm']
RIGHT = ['EjeiI','GA6KO','VlmQe','XSpDY','C7b7J','vH1ov','r1nOvb','DlCRn','xu45m']

def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def dump(path, data):
    Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

def validate(data, catalog):
    if data.get('version') != 1 or not data.get('operators'):
        raise ValueError('version must be 1; operators must be nonempty')
    seen=set()
    for job in data['operators']:
        name=job['name']
        if name in seen or name not in catalog:
            raise ValueError(f'Unknown or duplicate operator: {name}')
        seen.add(name)
        slots=job['slots']
        if len(slots)!=8 or {s['slot'] for s in slots}!=set(range(1,9)):
            raise ValueError(f'{name}: require exactly slots 1..8')
        for s in slots:
            if type(s['slot']) is not int or type(s.get('owned')) is not bool:
                raise ValueError(f'{name}: slot integer / owned boolean required')
            if not isinstance(s.get('player'),str) or not s['player'].strip():
                raise ValueError(f'{name}: player required')
            if type(s.get('visible',True)) is not bool:
                raise ValueError('visible must be boolean')
            vals=[s.get(k) for k in ('elite','level','potential')]
            if not s['owned']:
                if any(v is not None for v in vals):
                    raise ValueError(f'{name}/{s["slot"]}: unowned stats must be null')
                continue
            e,level,pot=vals
            if not all(type(v) is int for v in vals) or e not in (0,1,2) or not 1<=level<=90 or not 1<=pot<=6:
                raise ValueError(f'{name}/{s["slot"]}: elite 0..2, level 1..90, potential 1..6 integers required')
            asset=catalog[name].get(str(s['slot']))
            if asset is None:
                raise ValueError(f'{name}/{s["slot"]}: no configured source artwork; add catalog entry explicitly')
            for rel in [asset['url'],f'assets/psd_sources/elite{e}.png',f'assets/psd_sources/potential{pot}.png']:
                p=(ROOT/rel).resolve()
                if not p.is_relative_to(ROOT) or not p.is_file():
                    raise ValueError(f'Missing or out-of-project asset: {rel}')
            if sha(ROOT/asset['url'])!=asset['sha256']:
                raise ValueError(f'Artwork changed: {asset["url"]}')
    return data

def node(job,s,catalog):
    i=s['slot']-1;l=i<4;x=157 if l else 1180;y=YS[i];ids=LEFT if l else RIGHT
    owned=s['owned']
    d={ids[0]:{'enabled':owned},ids[1]:{'enabled':not owned},ids[2]:{'enabled':not owned},ids[4]:{'enabled':owned},ids[6]:{'enabled':owned},ids[8]:{'enabled':owned}}
    if owned:
        e=s['elite'];d[ids[3]]={'content':str(s['level'])}
        d[ids[5]]={'fill':{'type':'image','url':f'assets/psd_sources/elite{e}.png','mode':'stretch'},'x':(15 if l else 492) if e==2 else (12 if l else 489),'y':90 if e==2 else 97,'width':76 if e==2 else 81,'height':60 if e==2 else 49}
        d[ids[7]]={'fill':{'type':'image','url':f'assets/psd_sources/potential{s["potential"]}.png','mode':'fit'}}
        a=catalog[job['name']][str(s['slot'])];ax,ay,bx,by=a['bbox'];dx=ax-x;dy=ay-y;w=bx-ax;h=by-ay
        points=[(0,7.184),(536.526,7.184),(563.148,35.676),(455.548,176.896),(0,176.896)] if l else [(114.45,7.1),(584,7.1),(584,176.8),(33.5,176.8),(6.85,150.32)]
        d[ids[0]].update({'x':dx,'y':dy,'width':w,'height':h,'viewBox':[0,0,w,h],'geometry':'M'+' L'.join(f'{px-dx} {py-dy}' for px,py in points)+' Z','fill':{'type':'image','url':a['url'],'mode':'stretch'}})
    return {'type':'ref','ref':'HVZ4i' if l else 'wUVzi','name':f'{job["name"]}_SLOT_{s["slot"]:02d}_{s["player"]}','enabled':s.get('visible',True),'x':x,'y':y,'descendants':d}

def compile_batch(input_path,catalog_path,out):
    catalog=read(catalog_path);data=validate(read(input_path),catalog)
    out=Path(out).resolve()
    if out.exists():
        raise ValueError('Output directory already exists; choose a new run directory')
    out.mkdir(parents=True)
    run=sha(input_path)[:12];commands=[]
    for job in data['operators']:
        nodes=[node(job,s,catalog) for s in sorted(job['slots'],key=lambda s:s['slot'])]
        title=f'JSON_BATCH_{run}_{job["name"]}'
        export=(out/job['name']).as_posix()
        js='const nodes='+json.dumps(nodes,ensure_ascii=False)+';\n'
        js+='for(const n of nodes){Get(n.ref,{depth:0});for(const id of Object.keys(n.descendants))Get(id,{depth:0});}\n'
        js+=f'const matches=Get(n=>n.name==={json.dumps(title,ensure_ascii=False)}?n.id:undefined);if(matches.length>1)throw new Error("Duplicate batch frame");let frame=matches[0];\n'
        js+='if(!frame){const p=FindEmptySpace({width:1920,height:1080,padding:120,direction:"bottom"});frame=Insert(document,{type:"frame",name:'+json.dumps(title,ensure_ascii=False)+',x:p.x,y:p.y,width:1920,height:1080,layout:"none",clip:true,placeholder:true});for(const n of nodes){const id=Insert(document,n);Move(id,frame);}Update(frame,{placeholder:false});}\n'
        js+='const actual=Get(frame,{depth:1});if(actual.children.length!==8)throw new Error("Expected 8 slots");Print(actual);TakeScreenshot([frame]);Export([frame],"png",'+json.dumps(export,ensure_ascii=False)+',{scale:1});'
        p=out/f'{job["name"]}.pencil.js';p.write_text(js,encoding='utf-8')
        commands.append({'operator':job['name'],'filePath':str(ROOT/'character_cards_8slot.pen'),'inputFile':str(p),'outputDirectory':export,'frameName':title})
    dump(out/'run.json',{'input':str(Path(input_path).resolve()),'input_sha256':sha(input_path),'catalog_sha256':sha(catalog_path),'status':'prepared_not_rendered','commands':commands})
    print(f'Prepared {len(commands)} Pencil jobs: {out}/run.json')

def seed(manifest, output, catalog_path):
    old=read(manifest);data={'version':1,'operators':[]};catalog={}
    for j in old['jobs']:
        catalog[j['name']]={};slots=[]
        for s in j['slots']:
            slots.append({k:s[k] for k in ('slot','player','owned','elite','level','potential')})
            if s.get('override'):slots[-1]['note']=s['override']
            if s.get('art'):
                catalog[j['name']][str(s['slot'])]={**s['art'],'sha256':sha(ROOT/s['art']['url'])}
        data['operators'].append({'name':j['name'],'slots':slots})
    for p in (output,catalog_path):
        if Path(p).exists():raise ValueError(f'Refusing overwrite: {p}')
        Path(p).parent.mkdir(parents=True,exist_ok=True)
    dump(output,data);dump(catalog_path,catalog)

def main():
    parser=argparse.ArgumentParser(description=__doc__);sub=parser.add_subparsers(dest='command',required=True)
    p=sub.add_parser('seed');p.add_argument('--manifest',required=True);p.add_argument('--output',required=True);p.add_argument('--catalog',required=True)
    p=sub.add_parser('from-xlsx');p.add_argument('xlsx');p.add_argument('--output',required=True);p.add_argument('--overrides',help='JSON mapping cell address to explicitly approved replacement')
    for name in ('validate','prepare'):
        p=sub.add_parser(name);p.add_argument('input');p.add_argument('--catalog',default=str(ROOT/'config/cards_assets.json'))
        if name=='prepare':p.add_argument('--out',required=True)
    a=parser.parse_args()
    if a.command=='from-xlsx':
        import openpyxl
        from openpyxl.utils import get_column_letter
        w=openpyxl.load_workbook(a.xlsx,read_only=True,data_only=True);s=w['Sheet1']
        overrides=read(a.overrides) if a.overrides else {}
        result={'version':1,'operators':[]}
        for col in range(3,8):
            slots=[]
            for i in range(8):
                cells=[f'{get_column_letter(col)}{r+i}' for r in (4,13,22)]
                vals=[overrides.get(c,s[c].value) for c in cells]
                slots.append(dict(slot=i+1,player=s.cell(4+i,2).value,owned=any(v is not None for v in vals),elite=vals[0],level=vals[1],potential=vals[2]))
            result['operators'].append({'name':s.cell(2,col).value,'slots':slots})
        validate(result,read(ROOT/'config/cards_assets.json'))
        if Path(a.output).exists():raise ValueError('Refusing overwrite')
        Path(a.output).parent.mkdir(parents=True,exist_ok=True);dump(a.output,result)
    elif a.command=='seed':seed(a.manifest,a.output,a.catalog)
    elif a.command=='validate':validate(read(a.input),read(a.catalog));print('Input and assets valid')
    else:compile_batch(a.input,a.catalog,a.out)

if __name__=='__main__':
    main()
