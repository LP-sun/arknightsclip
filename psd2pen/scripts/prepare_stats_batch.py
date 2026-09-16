"""Read workbook and extract PSD pixels; emit Pencil commands, never write .pen."""
import json, hashlib
from pathlib import Path
import openpyxl
from psd_tools import PSDImage

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'delivery/stats_batch'
OUT.mkdir(parents=True,exist_ok=True)
ASSETS=ROOT/'assets/stats_batch'
ASSETS.mkdir(parents=True,exist_ok=True)
wb=openpyxl.load_workbook(ROOT/'统计.xlsx',data_only=True,read_only=True)
ws=wb['Sheet1']
audit=json.loads((ROOT/'reports/five_psd_template_audit.json').read_text(encoding='utf-8'))
jobs=[]
for col in range(3,8):
    name=ws.cell(2,col).value
    psd=PSDImage.open(ROOT/f'{name}.psd')
    group=next(l for l in psd if l.name=='卡面')
    source=next(a for a in audit['audits'] if a['file']==f'{name}.psd')
    slots=[]
    for idx in range(8):
        e,level,pot=[ws.cell(r+idx,col).value for r in (4,13,22)]
        override=None
        if name=='推王' and idx==4:
            e=2;override='User confirmed missing D8 as elite 2 from original PSD'
        owned=any(v is not None for v in (e,level,pot))
        if owned:
            assert e in (0,1,2) and isinstance(level,(int,float)) and pot in range(1,7),(name,idx,e,level,pot)
        art=None
        if owned:
            candidates=source['slots'][idx]['character_layers']
            usable=[x for x in candidates if x['effective_visible']]
            chosen=(usable or candidates)[0]
            layer=list(group)[chosen['order']]
            path=ASSETS/f'{col-2:02d}_slot{idx+1:02d}.png'
            layer.topil().save(path)
            art={'url':path.relative_to(ROOT).as_posix(),'bbox':list(layer.bbox),'source_layer':layer.name}
        from openpyxl.utils import get_column_letter
        slots.append({'slot':idx+1,'player':ws.cell(4+idx,2).value,'owned':owned,'elite':e,'level':level,'potential':pot,'cells':[f'{get_column_letter(col)}{r+idx}' for r in (4,13,22)],'override':override,'art':art})
    jobs.append({'name':name,'slots':slots,'psd_sha256':hashlib.sha256((ROOT/f'{name}.psd').read_bytes()).hexdigest()})
manifest={'source':'统计.xlsx','sheet':'Sheet1','range':'A1:G29','xlsx_sha256':hashlib.sha256((ROOT/'统计.xlsx').read_bytes()).hexdigest(),'jobs':jobs}
(OUT/'input_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
for job in jobs:
    code=['const p=FindEmptySpace({width:1920,height:1080,padding:120,direction:"bottom"});',f'const frame=Insert(document,{{type:"frame",name:{json.dumps("STATS_BATCH_"+job["name"],ensure_ascii=False)},x:p.x,y:p.y,width:1920,height:1080,layout:"none",clip:true,placeholder:true}});']
    for s in job['slots']:
        i=s['slot']-1;l=i<4;x=157 if l else 1180;y=([187,401,615,829,101,315,529,743][i]); owned=s['owned']
        ids=['DLOWp','B8gA3U','MxQrR','bFlti','QEG16','Q2ZirJ','coU5m','RB169','Q3ckUm'] if l else ['EjeiI','GA6KO','VlmQe','XSpDY','C7b7J','vH1ov','r1nOvb','DlCRn','xu45m']
        d={ids[0]:{'enabled':owned},ids[1]:{'enabled':not owned},ids[2]:{'enabled':not owned},ids[4]:{'enabled':owned},ids[6]:{'enabled':owned},ids[8]:{'enabled':owned}}
        if owned:
            e=s['elite'];d[ids[3]]={'content':str(int(s['level']))}
            d[ids[5]]={'fill':{'type':'image','url':f'assets/psd_sources/elite{e}.png','mode':'stretch'},'x':(15 if l else 492) if e==2 else (12 if l else 489),'y':90 if e==2 else 97,'width':76 if e==2 else 81,'height':60 if e==2 else 49}
            d[ids[7]]={'fill':{'type':'image','url':f'assets/psd_sources/potential{s["potential"]}.png','mode':'fit'}}
            a=s['art']; ax,ay,bx,by=a['bbox'];dx=ax-x;dy=ay-y;w=bx-ax;h=by-ay
            points=[(0,7.184),(536.526,7.184),(563.148,35.676),(455.548,176.896),(0,176.896)] if l else [(114.45,7.1),(584,7.1),(584,176.8),(33.5,176.8),(6.85,150.32)]
            geometry='M'+' L'.join(f'{px-dx} {py-dy}' for px,py in points)+' Z'
            d[ids[0]].update({'x':dx,'y':dy,'width':w,'height':h,'viewBox':[0,0,w,h],'geometry':geometry,'fill':{'type':'image','url':a['url'],'mode':'stretch'}})
        node={'type':'ref','ref':'HVZ4i' if l else 'wUVzi','name':f'STATS_{job["name"]}_SLOT_{i+1:02d}_{s["player"]}','x':x,'y':y,'descendants':d}
        code.append('const n'+str(i)+'=Insert(document,'+json.dumps(node,ensure_ascii=False)+');Move(n'+str(i)+',frame);')
    code+=['Update(frame,{placeholder:false});Print(frame);TakeScreenshot([frame]);',f'Export([frame],"png","E:/明日方舟报菜名/psd2pen/delivery/stats_batch/{job["name"]}",{{scale:1}});']
    (OUT/f'{job["name"]}.pencil.js').write_text('\n'.join(code),encoding='utf-8')
print(json.dumps({'jobs':[{'name':j['name'],'owned':sum(s['owned'] for s in j['slots'])} for j in jobs]},ensure_ascii=False))
